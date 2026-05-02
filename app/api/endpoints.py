from fastapi import APIRouter, UploadFile, File, HTTPException
from pydantic import BaseModel
from typing import Optional, Dict
import uuid
from app.services.file_parser import extract_text_from_file
from app.services.llm_service import analyze_resume_with_llm
import io
from fastapi.responses import StreamingResponse
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors

router = APIRouter()

# Simple in-memory storage to map IDs to extracted text
# (In a real production app, you would use a database like PostgreSQL or Redis)
document_storage: Dict[str, str] = {}

class AnalyzeRequest(BaseModel):
    resume_id: str
    jd_id: Optional[str] = None

class SaveDocumentRequest(BaseModel):
    text: str

class DownloadResumeRequest(BaseModel):
    full_rewritten_resume: str
    format: str = "docx"

@router.post("/upload-resume")
async def upload_resume(file: UploadFile = File(...)):
    text = await extract_text_from_file(file)
    resume_id = str(uuid.uuid4())
    document_storage[resume_id] = text
    return {
        "resume_id": resume_id,
        "text": text
    }

@router.post("/upload-jd")
async def upload_jd(file: UploadFile = File(...)):
    text = await extract_text_from_file(file)
    jd_id = str(uuid.uuid4())
    document_storage[jd_id] = text
    return {
        "jd_id": jd_id,
        "text": text
    }

@router.post("/save-document")
async def save_document(request: SaveDocumentRequest):
    doc_id = str(uuid.uuid4())
    document_storage[doc_id] = request.text
    return {"id": doc_id}

@router.post("/analyze")
async def analyze_resume(request: AnalyzeRequest):
    if request.resume_id not in document_storage:
        raise HTTPException(status_code=404, detail="Resume ID not found")
    resume_text = document_storage[request.resume_id]

    jd_text = ""
    if request.jd_id:
        if request.jd_id not in document_storage:
            raise HTTPException(status_code=404, detail="JD ID not found")
        jd_text = document_storage[request.jd_id]

    # Call LLM
    result = await analyze_resume_with_llm(resume_text, jd_text)
    return result

@router.post("/download-resume")
async def download_resume(request: DownloadResumeRequest):
    try:
        if request.format.lower() == "pdf":
            f = io.BytesIO()
            doc = SimpleDocTemplate(f, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
            styles = getSampleStyleSheet()
            
            title_style = ParagraphStyle('Title', parent=styles['Heading1'], alignment=TA_CENTER, fontSize=16, spaceAfter=2, fontName="Helvetica-Bold")
            contact_style = ParagraphStyle('Contact', parent=styles['Normal'], alignment=TA_CENTER, fontSize=10, spaceAfter=15, fontName="Helvetica")
            heading_style = ParagraphStyle('SectionHeading', parent=styles['Heading2'], fontSize=12, spaceBefore=12, spaceAfter=2, fontName="Helvetica-Bold", textTransform="uppercase")
            normal_style = ParagraphStyle('NormalText', parent=styles['Normal'], fontSize=10, spaceAfter=4, fontName="Helvetica")
            bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'], fontSize=10, leftIndent=15, spaceAfter=4, bulletIndent=5, fontName="Helvetica")
            
            story = []
            lines = request.full_rewritten_resume.split('\n')
            
            line_idx = 0
            while line_idx < len(lines):
                line = lines[line_idx].strip()
                if not line:
                    line_idx += 1
                    continue
                    
                if line_idx == 0 and not line.startswith('#'):
                    story.append(Paragraph(line, title_style))
                elif line_idx == 1 and not line.startswith('#'):
                    story.append(Paragraph(line, contact_style))
                elif line.startswith('#'):
                    text = line.replace('#', '').strip().upper()
                    story.append(Paragraph(text, heading_style))
                    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=8))
                elif line.startswith('-') or line.startswith('*'):
                    text = line[1:].strip()
                    story.append(Paragraph(f"• {text}", bullet_style))
                else:
                    if len(line) < 100 and '|' in line:
                        story.append(Paragraph(f"<b>{line}</b>", normal_style))
                    else:
                        story.append(Paragraph(line, normal_style))
                
                line_idx += 1
            
            doc.build(story)
            f.seek(0)
            
            return StreamingResponse(
                f,
                media_type="application/pdf",
                headers={"Content-Disposition": "attachment; filename=ATS_Optimized_Resume.pdf"}
            )
        else:
            doc = Document()
            
            lines = request.full_rewritten_resume.split('\n')
            for line in lines:
                if line.strip().startswith('#'):
                    level = line.count('#')
                    text = line.replace('#', '').strip()
                    doc.add_heading(text, level=min(level, 9))
                elif line.strip():
                    doc.add_paragraph(line.strip())

            f = io.BytesIO()
            doc.save(f)
            f.seek(0)

            return StreamingResponse(
                f,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": "attachment; filename=ATS_Optimized_Resume.docx"}
            )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
